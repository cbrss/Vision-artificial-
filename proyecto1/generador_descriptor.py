import cv2
import numpy as np
import os
import glob
from openpyxl import Workbook, load_workbook

# Generador de descriptores (invariantes de Hu)
# - Muestra solo contornos
# - Al presionar ESPACIO guarda los 7 invariantes de Hu en Excel

DATASET_XLSX = os.path.join(os.path.dirname(__file__), 'dataset_hu.xlsx')

# Carga de contornos de referencia desde clases/<categoria>/*.png|*.jpg|...
carpetas_referencia = {}
ruta_base = os.path.join(os.path.dirname(__file__), 'clases')

if os.path.exists(ruta_base):
    for item in os.listdir(ruta_base):
        ruta_completa = os.path.join(ruta_base, item)
        if os.path.isdir(ruta_completa):
            formatos_imagen = ["*.jpg", "*.jpeg", "*.png", "*.bmp"]
            imagenes = []
            for formato in formatos_imagen:
                patron_imagenes = os.path.join(ruta_completa, formato)
                imagenes.extend(glob.glob(patron_imagenes))
            if imagenes:
                carpetas_referencia[item] = ruta_completa
else:
    print(f"Error: No se encontró la carpeta {ruta_base}")

contornos_ref = {}
for nombre_categoria, ruta_carpeta in carpetas_referencia.items():
    contornos_ref[nombre_categoria] = []
    formatos_imagen = ["*.jpg", "*.jpeg", "*.png", "*.bmp"]
    imagenes = []
    for formato in formatos_imagen:
        patron_imagenes = os.path.join(ruta_carpeta, formato)
        imagenes.extend(glob.glob(patron_imagenes))
    for ruta_imagen in imagenes:
        img_ref = cv2.imread(ruta_imagen, 0)
        if img_ref is None:
            continue
        _, ref_bin = cv2.threshold(img_ref, 127, 255, cv2.THRESH_BINARY)
        conts, _ = cv2.findContours(ref_bin, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        if conts:
            areas = [cv2.contourArea(c) for c in conts]
            idx_max = int(np.argmax(areas))
            contornos_ref[nombre_categoria].append(conts[idx_max])


def asegurar_excel_con_encabezado(ruta_excel: str) -> None:
    if not os.path.exists(ruta_excel):
        wb = Workbook()
        ws = wb.active
        ws.title = 'data'
        ws.append(['hu1', 'hu2', 'hu3', 'hu4', 'hu5', 'hu6', 'hu7', 'label'])
        wb.save(ruta_excel)


def append_hu(ruta_excel: str, hu_vals: list, label: str) -> None:
    asegurar_excel_con_encabezado(ruta_excel)
    wb = load_workbook(ruta_excel)
    ws = wb['data']
    ws.append([float(hu_vals[0]), float(hu_vals[1]), float(hu_vals[2]),
               float(hu_vals[3]), float(hu_vals[4]), float(hu_vals[5]),
               float(hu_vals[6]), str(label)])
    wb.save(ruta_excel)


def escalar_contorno(contorno, area_objetivo):
    area_actual = cv2.contourArea(contorno)
    if area_actual == 0:
        return contorno.copy()
    factor = (area_objetivo / area_actual) ** 0.5
    M = cv2.moments(contorno)
    if M['m00'] == 0:
        return contorno.copy()
    cx = int(M['m10']/M['m00'])
    cy = int(M['m01']/M['m00'])
    contorno_centrado = contorno - [cx, cy]
    contorno_escalado = (contorno_centrado * factor).astype(np.int32) + [cx, cy]
    return contorno_escalado


def procesar(frame):
    # ROI: 3/4 inferiores de la imagen
    h, w = frame.shape[:2]
    y0 = h // 4
    roi = frame[y0:h, :]

    # Leo los valores de los controles
    umbral = cv2.getTrackbarPos('Umbral', 'Ajustes')
    auto = cv2.getTrackbarPos('Auto', 'Ajustes')
    kernel_size = cv2.getTrackbarPos('Kernel', 'Ajustes')
    area_min = cv2.getTrackbarPos('Area min', 'Ajustes')
    match_umbral = cv2.getTrackbarPos('Match max', 'Ajustes') / 1000.0

    # A gris y suavizado
    gray = cv2.cvtColor(roi, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray, (5, 5), 0)

    # Umbral
    if auto:
        _, binaria = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    else:
        _, binaria = cv2.threshold(gray, umbral, 255, cv2.THRESH_BINARY)

    # Morfología
    k = max(1, kernel_size)
    kernel = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (k, k))
    binaria = cv2.morphologyEx(binaria, cv2.MORPH_OPEN, kernel)
    binaria = cv2.morphologyEx(binaria, cv2.MORPH_CLOSE, kernel)

    # Contornos y salida
    contornos, _ = cv2.findContours(binaria, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    contornos_filtrados = [c for c in contornos if cv2.contourArea(c) >= area_min]

    salida = frame.copy()
    cv2.rectangle(salida, (0, y0), (w, h), (255, 0, 255), 2)

    # Clasificación por matchShapes con referencias
    resultados = []  # lista de (contorno, nombre, dist)
    for cnt in contornos_filtrados:
        area = cv2.contourArea(cnt)
        mejor_nombre = 'desconocido'
        mejor_dist = float('inf')
        for nombre_categoria, lista_contornos in contornos_ref.items():
            for ref in lista_contornos:
                ref_escalado = escalar_contorno(ref, area)
                dist = cv2.matchShapes(cnt, ref_escalado, cv2.CONTOURS_MATCH_I3, 0.0)
                if dist < mejor_dist:
                    mejor_dist = dist
                    mejor_nombre = nombre_categoria
        if match_umbral > 0 and mejor_dist > match_umbral:
            mejor_nombre = 'desconocido'
        resultados.append((cnt, mejor_nombre, mejor_dist))

    # Dibujo y texto
    for cnt, nombre, _ in resultados:
        cnt_shifted = cnt + np.array([0, y0])
        cv2.drawContours(salida, [cnt_shifted], -1, (0, 255, 0), 2)
        M = cv2.moments(cnt)
        if M['m00'] != 0:
            cx = int(M['m10']/M['m00'])
            cy = int(M['m01']/M['m00']) + y0
            cv2.putText(salida, nombre, (cx, cy), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 0), 2)

    return binaria, salida, resultados, y0


def calcular_hu(cnt) -> list:
    m = cv2.moments(cnt)
    hu = cv2.HuMoments(m).flatten()
    return [float(v) for v in hu]


# Ventanas y barras
cv2.namedWindow('Ajustes')
cv2.createTrackbar('Umbral', 'Ajustes', 127, 255, lambda x: None)
cv2.createTrackbar('Auto', 'Ajustes', 0, 1, lambda x: None)
cv2.createTrackbar('Kernel', 'Ajustes', 3, 20, lambda x: None)
cv2.createTrackbar('Area min', 'Ajustes', 100, 10000, lambda x: None)
cv2.createTrackbar('Match max', 'Ajustes', 20, 200, lambda x: None)  # 0.02 por defecto, máximo 0.20

# Captura de video
cap = cv2.VideoCapture(0)
if not cap.isOpened():
    raise Exception('No se pudo abrir la cámara')

print('Presiona ESPACIO para guardar Hu del contorno más grande. ESC para salir.')
while True:
    ret, frame = cap.read()
    if not ret:
        print('No se pudo leer el frame de la cámara')
        break

    binaria, salida, resultados, y0 = procesar(frame)

    # Mostrar
    cv2.imshow('Salida', salida)

    key = cv2.waitKey(10) & 0xFF
    if key == 27:  # ESC
        break
    if key == 32:  # ESPACIO
        if resultados:
            # Tomo el contorno con mayor área entre resultados
            idx = int(np.argmax([cv2.contourArea(r[0]) for r in resultados]))
            cnt, nombre, _ = resultados[idx]
            hu_vals = calcular_hu(cnt)
            append_hu(DATASET_XLSX, hu_vals, nombre)
            print('Hu guardados:', hu_vals, 'label:', nombre)
        else:
            print('No hay contornos válidos para guardar.')

cap.release()
cv2.destroyAllWindows()
