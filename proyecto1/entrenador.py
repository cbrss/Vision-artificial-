import os
import numpy as np
from openpyxl import load_workbook
from sklearn import tree
from sklearn.preprocessing import LabelEncoder
from joblib import dump


def cargar_dataset_desde_excel(ruta_excel: str):
    wb = load_workbook(ruta_excel)
    ws = wb['data']
    X = []
    Y = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None:
            continue
        vals = row[:7]
        label = row[7]
        # Saltar filas incompletas o sin etiqueta
        if None in vals or label in (None, ""):
            continue
        X.append([float(v) for v in vals])
        Y.append(label)
    return np.array(X, dtype=float), np.array(Y, dtype=object)


def main():
    base_dir = os.path.dirname(__file__)
    ruta_excel = os.path.join(base_dir, 'dataset_hu.xlsx')
    if not os.path.exists(ruta_excel):
        raise FileNotFoundError(f'No existe el dataset en {ruta_excel}. Genera muestras con clasificador.py')

    X, Y = cargar_dataset_desde_excel(ruta_excel)
    if X.size == 0 or Y.size == 0:
        raise ValueError('El dataset está vacío o sin etiquetas. Complete la columna label.')

    # Codificar etiquetas (admite texto o números)
    encoder = LabelEncoder()
    y_encoded = encoder.fit_transform(Y)

    clf = tree.DecisionTreeClassifier(random_state=42)
    clf.fit(X, y_encoded)

    # Guardar modelo
    modelo_path = os.path.join(base_dir, 'modelo_hu.joblib')
    artifact = {
        'model': clf,
        'classes': encoder.classes_.tolist(),
    }
    dump(artifact, modelo_path)
    print(f'Modelo entrenado y guardado en {modelo_path}. Clases: {encoder.classes_.tolist()}')


if __name__ == '__main__':
    main()


